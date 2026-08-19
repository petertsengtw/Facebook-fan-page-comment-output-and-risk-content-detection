"""分類品質評估測試（規格 §8.4）。"""
import pytest

from smc.classify.evaluate import (
    compute_metrics,
    confusion_pairs,
    load_goldset,
    macro_f1,
    write_goldset_template,
)
from smc.classify.labels import Label

ATTACK = Label.PERSONAL_ATTACK
ORG = Label.ORG_ATTACK
NEUTRAL = Label.NEUTRAL


def test_perfect_agreement():
    gold = {"1": {ATTACK}, "2": {NEUTRAL}}
    assert compute_metrics(gold, gold)[ATTACK].f1 == 1.0


def test_missed_label_lowers_recall_not_precision():
    """漏報：人工說有、系統沒抓到 → recall 降低，precision 不受影響。"""
    m = compute_metrics({"1": {ATTACK}, "2": {ATTACK}}, {"1": {ATTACK}, "2": set()})[ATTACK]
    assert m.recall == 0.5
    assert m.precision == 1.0
    assert m.support == 2


def test_false_positive_lowers_precision_not_recall():
    m = compute_metrics({"1": {ATTACK}, "2": set()}, {"1": {ATTACK}, "2": {ATTACK}})[ATTACK]
    assert m.precision == 0.5
    assert m.recall == 1.0


def test_zero_support_does_not_divide_by_zero():
    m = compute_metrics({"1": set()}, {"1": {ATTACK}})[ATTACK]
    assert (m.precision, m.recall, m.f1) == (0.0, 0.0, 0.0)


def test_only_shared_comments_are_evaluated():
    """標註集中尚未分類的留言不得影響指標。"""
    metrics = compute_metrics({"1": {ATTACK}, "未分類": {ATTACK}}, {"1": {ATTACK}})
    assert metrics[ATTACK].support == 1


def test_macro_f1_is_not_dominated_by_majority_label():
    """規格 §8.4：用 macro 而非 micro，否則 neutral 會稀釋掉嚴重問題。

    此例中 neutral 完全正確（99 則），personal_attack 完全失敗（1 則）。
    macro F1 應明顯低於 1，讓失敗可見。
    """
    gold = {str(i): {NEUTRAL} for i in range(99)} | {"attack": {ATTACK}}
    pred = {str(i): {NEUTRAL} for i in range(99)} | {"attack": set()}
    assert macro_f1(compute_metrics(gold, pred)) < 0.6


def test_confusion_pairs_report_actual_to_wrong():
    gold = {"1": {ORG}, "2": {ORG}}
    pred = {"1": {ATTACK}, "2": {ATTACK}}
    assert confusion_pairs(gold, pred)[0] == (ORG, ATTACK, 2)


# ── goldset 檔案處理 ──────────────────────────────────────────────

def test_goldset_round_trip(tmp_path):
    path = tmp_path / "g.csv"
    write_goldset_template(path, [
        {"comment_id": "c1", "message": "測試", "labels": "org_attack;sarcasm", "note": ""},
        {"comment_id": "c2", "message": "普通", "labels": "", "note": ""},
    ])
    gold = load_goldset(path)
    assert gold["c1"] == {Label.ORG_ATTACK, Label.SARCASM}
    assert gold["c2"] == set(), "空標籤代表「無任何標籤」，不是未標註"


def test_goldset_rejects_unknown_label(tmp_path):
    path = tmp_path / "g.csv"
    write_goldset_template(path, [
        {"comment_id": "c1", "message": "x", "labels": "typo_label", "note": ""},
    ])
    with pytest.raises(ValueError, match="不在標籤集中"):
        load_goldset(path)


def test_goldset_rejects_missing_columns(tmp_path):
    path = tmp_path / "g.csv"
    path.write_text("comment_id,message\nc1,x\n", encoding="utf-8")
    with pytest.raises(ValueError, match="缺少欄位"):
        load_goldset(path)


def test_goldset_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_goldset(tmp_path / "nope.csv")


# ── Excel 標註表 ──────────────────────────────────────────────────

def test_xlsx_goldset_round_trip(tmp_path):
    from smc.classify.evaluate import load_any_goldset, write_goldset_xlsx

    path = tmp_path / "g.xlsx"
    write_goldset_xlsx(path, [
        {"comment_id": "c1", "message": "測試留言", "note": ""},
        {"comment_id": "c2", "message": "另一則", "note": "已隱藏"},
    ])
    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["標註", "標籤說明"]
    ws = wb["標註"]
    ws["C2"] = "org_attack"
    ws["D2"] = "sarcasm"
    wb.save(path)

    gold = load_any_goldset(path)
    assert gold["c1"] == {ORG, Label.SARCASM}
    assert gold["c2"] == set(), "留空代表無標籤，非未標註"


def test_xlsx_rejects_unknown_label(tmp_path):
    from openpyxl import load_workbook

    from smc.classify.evaluate import load_any_goldset, write_goldset_xlsx

    path = tmp_path / "g.xlsx"
    write_goldset_xlsx(path, [{"comment_id": "c1", "message": "x", "note": ""}])
    wb = load_workbook(path)
    wb["標註"]["C2"] = "手打錯的標籤"
    wb.save(path)
    with pytest.raises(ValueError, match="不在標籤集中"):
        load_any_goldset(path)


def test_negative_examples_count_toward_precision():
    """人工判定「沒問題」的留言必須納入評估。

    若只評估有標籤者，系統把乾淨留言誤標時不會計為誤報，precision 會被高估。
    """
    gold = {"1": {ATTACK}, "2": set(), "3": set()}
    pred = {"1": {ATTACK}, "2": {ATTACK}, "3": set()}
    m = compute_metrics(gold, pred)[ATTACK]
    assert m.false_positive == 1, "誤標乾淨留言須計為誤報"
    assert m.precision == 0.5


def test_coarse_metrics_collapse_system_labels():
    """人工的顆粒度較粗：不區分攻擊對象是機構還是個人。"""
    from smc.classify.evaluate import compute_coarse_metrics, to_coarse

    assert to_coarse({Label.ORG_ATTACK}) == {"謾罵／人身攻擊"}
    assert to_coarse({Label.PERSONAL_ATTACK}) == {"謾罵／人身攻擊"}
    assert to_coarse({Label.NEUTRAL}) == set()

    gold = {"1": {"謾罵／人身攻擊"}, "2": {"謾罵／人身攻擊"}}
    pred = {"1": {Label.ORG_ATTACK}, "2": {Label.PERSONAL_ATTACK}}
    assert compute_coarse_metrics(gold, pred)["謾罵／人身攻擊"].recall == 1.0


def test_detection_recall_ignores_label_disagreement():
    """標籤歧異與完全漏抓的後果天差地遠：前者小編照樣看得到。

    偵測率只問「有沒有被標出來」，不問標成什麼。
    """
    from smc.classify.evaluate import detection_recall

    gold = {"a": {"謾罵／人身攻擊"}, "b": {"謾罵／人身攻擊"}, "c": set()}
    severities = {"a": 2, "b": 0, "c": 3}
    result = detection_recall(gold, severities, threshold=2)
    assert result.total == 2, "只計人工認定有問題者"
    assert result.flagged == 1
    assert result.missed == ["b"]
    assert result.rate == 0.5
