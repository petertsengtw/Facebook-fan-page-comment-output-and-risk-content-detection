"""分類品質評估（規格 §8.4）。

多標籤任務，故各標籤獨立以二元分類方式計算 precision / recall / F1。
不使用整體 accuracy —— 在標籤極度不平衡的資料上（本專案 neutral 佔多數），
accuracy 會被多數類別灌水到失去意義。
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from .labels import CHINESE_NAME, Label

#: 人工複核常用的粗分類 → 本系統標籤（規格 §8.4）
#:
#: 人工判讀的顆粒度天然比系統粗：複核者關心「這是不是謾罵」，
#: 不需要區分攻擊對象是機構還是個人。強行要求人工標到系統的顆粒度
#: 只會製造假性不一致，故評估時將系統標籤收斂到人工的層級再比對。
COARSE_MAP: dict[str, set[Label]] = {
    "謾罵／人身攻擊": {Label.ORG_ATTACK, Label.PERSONAL_ATTACK, Label.THREAT},
    "嘲諷／陰陽怪氣": {Label.SARCASM},
    "負面質疑": {Label.COMPLAINT, Label.MISINFORMATION},
    "洗版／離題": {Label.SPAM},
}
FINE_TO_COARSE: dict[Label, str] = {
    fine: coarse for coarse, fines in COARSE_MAP.items() for fine in fines
}


def to_coarse(labels: set[Label]) -> set[str]:
    """將系統標籤收斂為人工複核的粗分類。"""
    return {FINE_TO_COARSE[x] for x in labels if x in FINE_TO_COARSE}


def is_coarse_goldset(gold: dict[str, set]) -> bool:
    """判斷標註集使用的是粗分類還是系統標籤。"""
    return any(
        isinstance(v, str) and v in COARSE_MAP
        for labels in gold.values()
        for v in labels
    )


def compute_coarse_metrics(
    gold: dict[str, set[str]], predicted: dict[str, set[Label]]
) -> dict[str, LabelMetrics]:
    """以人工的粗分類為基準計算指標。"""
    coarse_pred = {cid: to_coarse(labs) for cid, labs in predicted.items()}
    shared = gold.keys() & coarse_pred.keys()
    names = {n for cid in shared for n in gold[cid] | coarse_pred[cid]}
    out: dict[str, LabelMetrics] = {}
    for name in sorted(names):
        tp = sum(1 for c in shared if name in gold[c] and name in coarse_pred[c])
        fp = sum(1 for c in shared if name not in gold[c] and name in coarse_pred[c])
        fn = sum(1 for c in shared if name in gold[c] and name not in coarse_pred[c])
        out[name] = LabelMetrics(name, tp, fp, fn)  # type: ignore[arg-type]
    return out


#: goldset CSV 的必要欄位
GOLDSET_FIELDS = ("comment_id", "message", "labels", "note")


@dataclass(frozen=True)
class LabelMetrics:
    label: Label
    true_positive: int
    false_positive: int
    false_negative: int

    @property
    def support(self) -> int:
        """人工標註中該標籤的實際則數。"""
        return self.true_positive + self.false_negative

    @property
    def predicted(self) -> int:
        return self.true_positive + self.false_positive

    @property
    def precision(self) -> float:
        return self.true_positive / self.predicted if self.predicted else 0.0

    @property
    def recall(self) -> float:
        return self.true_positive / self.support if self.support else 0.0

    @property
    def f1(self) -> float:
        p, r = self.precision, self.recall
        return 2 * p * r / (p + r) if (p + r) else 0.0

    @property
    def chinese_name(self) -> str:
        if isinstance(self.label, str) and not isinstance(self.label, Label):
            return self.label
        return CHINESE_NAME.get(self.label, str(self.label))


def compute_metrics(
    gold: dict[str, set[Label]], predicted: dict[str, set[Label]]
) -> dict[Label, LabelMetrics]:
    """計算各標籤指標。僅納入兩邊都有資料的留言。"""
    shared = gold.keys() & predicted.keys()
    labels = {lab for cid in shared for lab in gold[cid] | predicted[cid]}
    metrics = {}
    for label in sorted(labels):
        tp = sum(1 for cid in shared if label in gold[cid] and label in predicted[cid])
        fp = sum(1 for cid in shared if label not in gold[cid] and label in predicted[cid])
        fn = sum(1 for cid in shared if label in gold[cid] and label not in predicted[cid])
        metrics[label] = LabelMetrics(label, tp, fp, fn)
    return metrics


def macro_f1(metrics: dict[Label, LabelMetrics]) -> float:
    """各標籤 F1 的未加權平均。

    用 macro 而非 micro —— micro 會被 neutral 這類多數標籤主導，
    使得「人身攻擊完全抓不到」這種嚴重問題被稀釋成小數點後的差異。
    """
    if not metrics:
        return 0.0
    return sum(m.f1 for m in metrics.values()) / len(metrics)


def confusion_pairs(
    gold: dict[str, set[Label]], predicted: dict[str, set[Label]], *, top: int = 10
) -> list[tuple[Label, Label, int]]:
    """最常見的混淆組合：(人工標註, 系統誤判為, 次數)。"""
    from collections import Counter

    pairs: Counter = Counter()
    for cid in gold.keys() & predicted.keys():
        missed = gold[cid] - predicted[cid]
        extra = predicted[cid] - gold[cid]
        for actual in missed:
            for wrong in extra:
                pairs[(actual, wrong)] += 1
    return [(a, b, n) for (a, b), n in pairs.most_common(top)]


@dataclass(frozen=True)
class DetectionResult:
    """偵測率：不論標籤為何，人工認定有問題者是否被送到人工面前。

    標籤歧異與完全漏抓在指標上都會壓低 recall，但實務後果天差地遠：
    前者小編照樣看得到，後者則永遠不會被發現。此指標只問「有沒有被標出來」。
    """

    flagged: int
    total: int
    missed: list[str]

    @property
    def rate(self) -> float:
        return self.flagged / self.total if self.total else 0.0


def detection_recall(
    gold: dict[str, set], severities: dict[str, int], *, threshold: int = 2
) -> DetectionResult:
    positives = [cid for cid, labs in gold.items() if labs and cid in severities]
    flagged = [cid for cid in positives if severities[cid] >= threshold]
    missed = [cid for cid in positives if severities[cid] < threshold]
    return DetectionResult(len(flagged), len(positives), missed)


def load_goldset(path: Path) -> dict[str, set[Label]]:
    """讀取人工標註集。labels 欄以分號分隔，可為空（表示該則無任何標籤）。"""
    if not path.exists():
        raise FileNotFoundError(f"找不到標註集 {path}")
    result: dict[str, set[Label]] = {}
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        missing = set(GOLDSET_FIELDS[:3]) - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"標註集缺少欄位：{sorted(missing)}")
        for line_no, row in enumerate(reader, start=2):
            cid = (row.get("comment_id") or "").strip()
            if not cid:
                continue
            labels = set()
            for raw in (row.get("labels") or "").split(";"):
                if not (name := raw.strip()):
                    continue
                try:
                    labels.add(Label(name))
                except ValueError as exc:
                    raise ValueError(
                        f"第 {line_no} 行的標籤 {name!r} 不在標籤集中"
                    ) from exc
            result[cid] = labels
    return result


#: Excel 標註表的欄位
XLSX_HEADERS = ["comment_id", "留言內容", "標籤1", "標籤2", "標籤3", "備註"]
#: 供人工挑選的標籤（不含規則專用者）
SELECTABLE = [x for x in Label if x not in (Label.NO_TEXT, Label.CLASSIFY_ERROR)]


def write_goldset_xlsx(path: Path, rows: list[dict[str, str]]) -> Path:
    """輸出附下拉選單與標籤說明的 Excel 標註表。

    CSV 需要人工手打標籤名稱，容易拼錯且看不到定義；
    下拉選單可從源頭杜絕拼字錯誤。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "標註"
    ws.append(XLSX_HEADERS)
    for row in rows:
        existing = [x.strip() for x in (row.get("labels") or "").split(";") if x.strip()]
        existing += [""] * (3 - len(existing))
        ws.append([row["comment_id"], row["message"], *existing[:3], row.get("note", "")])

    # 下拉同時提供人工粗分類與系統標籤 —— 匯入的既有標註用前者
    options = '"' + ",".join([*COARSE_MAP, *(str(x) for x in SELECTABLE)]) + '"'
    dv = DataValidation(type="list", formula1=options, allow_blank=True, showDropDown=False)
    dv.error = "請由下拉選單選擇，或留空表示中性"
    dv.errorTitle = "標籤不在清單中"
    ws.add_data_validation(dv)
    if rows:
        for col in ("C", "D", "E"):
            dv.add(f"{col}2:{col}{len(rows) + 1}")

    header_fill = PatternFill("solid", fgColor="D9E2F3")
    for i in range(1, len(XLSX_HEADERS) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for i, width in enumerate([26, 74, 20, 20, 20, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"
    if rows:
        ws.auto_filter.ref = f"A1:F{len(rows) + 1}"

    guide = wb.create_sheet("標籤說明")
    guide.append(["標籤", "中文", "說明"])
    notes = {
        Label.PERSONAL_ATTACK: "針對**特定個人**（醫師、護理師、員工、其他留言者）的侮辱貶低",
        Label.ORG_ATTACK: "針對本院／慈濟**整體**的侮辱性字眼（詐騙集團、邪魔教、無恥）",
        Label.SARCASM: "反諷、酸言，未達謾罵程度",
        Label.POLITICS: "涉及政黨、選舉、政治人物（中性標記，不代表有問題）",
        Label.COMPLAINT: "對流程、等候、費用、醫療結果的具體不滿；**論理性批評也算這類**",
        Label.MISINFORMATION: "對醫療事實或本院措施的明顯錯誤陳述",
        Label.SPAM: "商品推銷、外部導流、無關重複張貼",
        Label.PII_LEAK: "出現身分證、病歷號、電話，或將個人姓名與病情連結",
        Label.SELF_HARM: "表達自殺或自傷意念（病人安全事項）",
        Label.THREAT: "暴力威脅、恐嚇（揚言提告不算）",
        Label.POSITIVE: "感謝、讚美、祝福",
        Label.NEUTRAL: "詢問、事實陳述、與風險無關",
    }
    for label in SELECTABLE:
        guide.append([str(label), CHINESE_NAME[label], notes.get(label, "")])

    guide.append([])
    guide.append(["最容易搞混的界線", "", ""])
    guide.append(["「慈濟善款不透明，帳目該公開」", "→ complaint", "論理性批評"])
    guide.append(["「慈濟是詐騙集團」", "→ org_attack", "侮辱性字眼"])
    guide.append(["差別在有無侮辱性字眼，不在立場強弱", "", ""])
    guide.append([])
    guide.append(["留空代表什麼", "", ""])
    guide.append(["三個標籤欄都留空 = 該則無任何標籤（中性）", "", ""])
    guide.append(["若不確定如何標，請填備註欄說明，先跳過", "", ""])

    for i in range(1, 4):
        guide.cell(row=1, column=i).font = Font(bold=True)
        guide.cell(row=1, column=i).fill = header_fill
    for i, width in enumerate([34, 16, 62], 1):
        guide.column_dimensions[get_column_letter(i)].width = width
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def load_goldset_xlsx(path: Path) -> dict[str, set[Label]]:
    """讀取 Excel 標註表。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["標註"] if "標註" in wb.sheetnames else wb.active
    result: dict[str, set[Label]] = {}
    for line_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        cid = str(row[0]).strip()
        labels = set()
        for raw in row[2:5]:
            if not (name := str(raw or "").strip()):
                continue
            if name in COARSE_MAP:
                labels.add(name)  # 人工粗分類，評估時再收斂系統標籤比對
                continue
            try:
                labels.add(Label(name))
            except ValueError as exc:
                raise ValueError(f"第 {line_no} 列的標籤 {name!r} 不在標籤集中") from exc
        result[cid] = labels
    return result


def load_any_goldset(path: Path) -> dict[str, set[Label]]:
    """依副檔名讀取標註集（.xlsx 或 .csv）。"""
    if not path.exists():
        raise FileNotFoundError(f"找不到標註集 {path}")
    return load_goldset_xlsx(path) if path.suffix.lower() == ".xlsx" else load_goldset(path)


def write_goldset_template(path: Path, rows: list[dict[str, str]]) -> Path:
    """輸出待人工標註的 CSV 範本。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=GOLDSET_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    return path


def import_manual_annotations(
    source: Path,
    *,
    lookup: dict[str, str],
    sheet: str = "留言總表",
    content_column: str = "留言內容",
    label_column: str = "標記",
    level_column: str = "層級",
) -> tuple[list[dict[str, str]], dict[str, int]]:
    """匯入既有的人工標註表（規格 §8.4）。

    以留言內容比對本系統的資料庫。三段式比對，因人工擷取的文字與 API 取得者
    常有差異：長留言被節錄、回覆前面接了父留言者姓名、空白與換行不同。

    回傳 (goldset 列, 統計)。
    """
    from openpyxl import load_workbook

    from smc.util.hashing import normalize_message

    wb = load_workbook(source, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"找不到工作表「{sheet}」，可用：{wb.sheetnames}")
    ws = wb[sheet]

    header = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, values_only=True))]
    for name in (content_column, label_column):
        if name not in header:
            raise ValueError(f"工作表缺少欄位「{name}」，實際欄位：{header}")
    ci, li = header.index(content_column), header.index(label_column)
    vi = header.index(level_column) if level_column in header else None

    keys = list(lookup)
    stats: dict[str, int] = {
        "已標註": 0, "無標記": 0, "圖片或空白": 0, "比對成功": 0, "找不到": 0
    }

    def find(text: str, is_reply: bool) -> str | None:
        key = normalize_message(text)
        if key in lookup:
            return lookup[key]
        # 回覆在人工擷取中前綴了父留言者姓名，去掉第一個詞再試
        if is_reply and " " in key:
            tail = key.split(" ", 1)[1]
            if tail in lookup:
                return lookup[tail]
            for k in keys:
                if k.startswith(tail[:20]) and len(tail) >= 8:
                    return lookup[k]
        # 長留言被節錄：以前綴比對
        if len(key) >= 12:
            for k in keys:
                if k.startswith(key[:20]):
                    return lookup[k]
        return None

    rows: list[dict[str, str]] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        content = str(raw[ci] or "").strip()
        label = str(raw[li] or "").strip()
        is_reply = bool(vi is not None and str(raw[vi] or "").strip() == "回覆")
        stats["已標註" if label else "無標記"] += 1
        if not content:
            stats["圖片或空白"] += 1
            continue
        cid = find(content, is_reply)
        if cid is None:
            stats["找不到"] += 1
            continue
        stats["比對成功"] += 1
        rows.append({
            "comment_id": cid,
            "message": content[:200],
            "labels": label,
            "note": "人工標註匯入",
        })
    return rows, stats
