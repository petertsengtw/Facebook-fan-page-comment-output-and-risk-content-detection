"""Excel 報表產生（規格 §9、附錄 B）。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from smc.classify.labels import CHINESE_NAME, INFORMATIONAL_ONLY, Label
from smc.classify.merge import hits_from_json
from smc.storage.repo import ReportRow

#: 嚴重度底色（規格 附錄 B）
SEVERITY_FILL = {
    3: PatternFill("solid", fgColor="FFC7CE"),  # 紅
    2: PatternFill("solid", fgColor="FFE0B2"),  # 橘
    1: PatternFill("solid", fgColor="FFF2CC"),  # 黃
}
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
SEVERITY_NAME = {0: "0 無需處理", 1: "1 留意", 2: "2 建議人工回應", 3: "3 需主管即刻處理"}

FOCUS_HEADERS = [
    "嚴重度", "標籤", "留言時間", "貼文標題", "貼文發布日", "層級",
    "留言內容", "判斷理由", "同內容則數", "留言連結", "comment_id",
    "已隱藏", "複核狀態", "複核備註",
]
FOCUS_WIDTHS = [10, 18, 17, 30, 12, 8, 55, 30, 11, 46, 24, 8, 12, 24]

ALL_HEADERS = [
    "嚴重度", "標籤", "留言時間", "貼文標題", "層級", "留言內容",
    "判斷理由", "留言連結", "已隱藏",
]
ALL_WIDTHS = [10, 18, 17, 30, 8, 55, 30, 46, 8]


def _labels_text(labels_json: str | None) -> str:
    if not labels_json:
        return "（未分類）"
    hits = hits_from_json(labels_json)
    if not hits:
        return "—"
    parts = []
    for hit in sorted(hits, key=lambda h: -h.severity):
        name = CHINESE_NAME.get(hit.label, str(hit.label))
        mark = "（中性標記）" if hit.label in INFORMATIONAL_ONLY else f"({hit.severity})"
        parts.append(f"{name}{mark}")
    return "、".join(parts)


def _style_sheet(ws: Worksheet, widths: list[int]) -> None:
    for i, _ in enumerate(widths, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def build_workbook(
    *,
    focus: list[ReportRow],
    everything: list[ReportRow],
    duplicates: list[tuple[str, int, int, str, list[str]]],
    label_counts: dict[str, int],
    severity_counts: dict[int, int],
    top_posts: list[tuple[str, str, int, int]],
    meta: dict[str, str],
    gaps: list[tuple[datetime, datetime]],
    dup_counts: dict[str, int],
) -> Workbook:
    wb = Workbook()

    # ── 工作表 1：需關注留言 ──
    ws = wb.active
    ws.title = "需關注留言"
    ws.append(FOCUS_HEADERS)
    for r in focus:
        ws.append([
            SEVERITY_NAME.get(r.max_severity, r.max_severity),
            _labels_text(r.labels_json),
            r.created_time.strftime("%Y-%m-%d %H:%M"),
            r.post_title,
            r.post_created.strftime("%Y-%m-%d"),
            "回覆" if r.is_reply else "主留言",
            r.message or "（純圖片／貼圖）",
            r.rationale or "",
            dup_counts.get(r.message_hash, 1),
            r.permalink_url or "",
            r.comment_id,
            "是" if r.is_hidden else "",
            r.review_status or "待複核",
            r.review_note or "",
        ])
        if fill := SEVERITY_FILL.get(r.max_severity):
            ws.cell(row=ws.max_row, column=1).fill = fill
    _style_sheet(ws, FOCUS_WIDTHS)

    # ── 工作表 2：全量留言 ──
    ws2 = wb.create_sheet("全量留言")
    ws2.append(ALL_HEADERS)
    for r in everything:
        ws2.append([
            SEVERITY_NAME.get(r.max_severity, r.max_severity),
            _labels_text(r.labels_json),
            r.created_time.strftime("%Y-%m-%d %H:%M"),
            r.post_title,
            "回覆" if r.is_reply else "主留言",
            r.message or "（純圖片／貼圖）",
            r.rationale or "",
            r.permalink_url or "",
            "是" if r.is_hidden else "",
        ])
    _style_sheet(ws2, ALL_WIDTHS)

    # ── 工作表 3：重複內容彙總（法務蒐證用）──
    ws3 = wb.create_sheet("重複內容彙總")
    ws3.append(["出現次數", "橫跨貼文數", "內容", "全部留言連結"])
    for _digest, count, posts, sample, links in duplicates:
        ws3.append([count, posts, sample, "\n".join(links)])
    _style_sheet(ws3, [10, 12, 60, 60])

    # ── 工作表 4：摘要統計 ──
    ws4 = wb.create_sheet("摘要統計")
    ws4.append(["項目", "值"])
    for key, value in meta.items():
        ws4.append([key, value])

    ws4.append([])
    ws4.append(["嚴重度分布", ""])
    for level in sorted(severity_counts, reverse=True):
        ws4.append([SEVERITY_NAME.get(level, str(level)), severity_counts[level]])

    ws4.append([])
    ws4.append(["標籤命中數", ""])
    for name, count in sorted(label_counts.items(), key=lambda kv: -kv[1]):
        try:
            zh = CHINESE_NAME[Label(name)]
        except ValueError:
            zh = name
        ws4.append([f"{zh}（{name}）", count])

    ws4.append([])
    ws4.append(["風險留言最多的貼文", ""])
    ws4.append(["貼文標題", "需關注則數"])
    for _pid, title, n, _sev in top_posts:
        ws4.append([title, n])

    if gaps:
        ws4.append([])
        ws4.append(["⚠ 未擷取區間（結果不可信）", ""])
        for start, end in gaps:
            ws4.append([f"{start:%Y-%m-%d} ～ {end:%Y-%m-%d}", "從未執行 fetch"])

    ws4.append([])
    ws4.append(["說明", "「政治相關」為中性資訊標記，不計入嚴重度，不得作為處置依據（規格 §5.5）"])
    ws4.append(["", "「自傷風險」屬病人安全事項，須優先關懷轉介，非負面輿情（規格 §5.7）"])
    ws4.append(["", "本報表不含留言者姓名 —— 標準存取下 API 不回傳；請經留言連結至 Facebook 查看"])
    blind_spot = (
        "⚠ 已知盲區：純圖片留言未經判讀。"
        "已有留言者將攻擊文字製成圖片以規避偵測（規格 §16 R10）"
    )
    ws4.append(["", blind_spot])
    _style_sheet(ws4, [46, 60])

    return wb


def save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
